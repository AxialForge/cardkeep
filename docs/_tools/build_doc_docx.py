"""Render a Markdown doc (headings, paragraphs, lists, tables, code) to a portrait Word file.

Usage: python docs/_tools/build_doc_docx.py docs/phone-tests/0.6.2.md "Documentation/CardKeep 0.6.2 Phone Test Plan Rev 1.0.docx"
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt, RGBColor

INK = RGBColor(0x0E, 0x1A, 0x2B)
BRASS = RGBColor(0x7A, 0x5B, 0x22)


def runs(par, text):
    # **bold** and `code`
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
        else:
            par.add_run(part)


def table(doc, rows):
    cells = [[c.replace("<br>", "").strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j, c in enumerate(r[: len(cells[0])]):
            p = t.cell(i, j).paragraphs[0]
            runs(p, c)
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
            if i > 0 and not c and len(cells[0]) == 2:
                p.add_run("\n\n")  # room to write notes
    doc.add_paragraph()


def main(src, out):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Mm(215.9), Mm(279.4)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Mm(16))
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = Pt(10)
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(8)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            table(doc, rows)
            continue
        elif ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            h = doc.add_heading(ln.lstrip("#").strip(), level=min(level, 3))
            for r in h.runs:
                r.font.color.rgb = INK if level < 3 else BRASS
        elif ln.strip() == "---":
            doc.add_page_break()
        elif re.match(r"^\d+\. ", ln):
            runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
        elif ln.startswith(("- ", "* ")):
            runs(doc.add_paragraph(style="List Bullet"), ln[2:])
        elif ln.strip():
            text = ln
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(#|\||-|\*|\d+\.|```)", lines[i + 1]):
                i += 1
                text += " " + lines[i].strip()
            runs(doc.add_paragraph(), text)
        i += 1
    doc.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
